"""
occ_tune_core.py
================
Baseline A / B / C 的 OCC 超參數調整：共用核心程式
（筆電測試版與 Azure 全跑版都 import 這支，兩邊只差 run_laptop.py / run_azure.py 的設定）

  A : 原始特徵 OF_maj              → OCC
  B : AE 深度特徵 DF_maj            → OCC
  C : [OF_maj , DF_maj] 串接        → OCC
  （B 和 C 用的是同一顆 AE，所以一格 AE 訓練一次，B、C 兩邊的結果一起算）

────────────────────────────────────────────────────────────────────────────
這支程式解決的三件事（對應 N 程式的三個問題）
────────────────────────────────────────────────────────────────────────────
 1. 存檔順序：每一「格」(資料集 × fold × AE × config) 算完，先寫進暫存檔，
    檢查列數正確後才 os.replace() 改成正式檔名。正式檔存在 ＝ 這格完成。
    不另外維護「進度表」，所以不可能出現「打了勾但答案沒寫進去」。
 2. 完成判定：一格裡面所有 OCC、所有參數組合都「成功」、列數也對，才寫正式檔。
    只要有任何一組參數失敗，整格不存檔、記進 errors.log，下次重跑會再補。
 3. 特徵快取：AE 抽出的 DF 存成 .npz，下次（或換一組 OCC 參數格）直接讀，
    不重訓 AE。快取帶有 AE 設定與資料檔指紋，不相符就自動重訓。

其他設計
  • OCC 參數迴圈在最內層：同一份特徵試完全部參數，AE 只訓練一次。
  • 每格固定亂數種子（由 資料集/fold/AE/config 算出），跟執行順序、
    幾個 worker 平行都無關 → 同一台機器上，中斷重跑與一次跑完的結果相同（已測試）。
    不同機器／套件版本之間不保證逐位元相同，所以正式結果只用一台機器（Azure）產生。
  • 全程只用 CPU（刻意關掉 GPU）：這麼小的網路 GPU 反而比較慢。
  • 開跑前嚴格檢查資料：指定的資料集必須一個不少，每個 5 折的 train/test 都要在、
    都讀得進來，否則直接停止（不會默默少跑一個資料集）。
  • 同一個 run 不能同時開兩個程式（作業系統檔案鎖，程式結束或當掉會自動釋放）。
  • 每個參數組合的結果全部保存，事後要用哪種方式挑參數都只要重算，不用重跑。
  • 設定簽章：會改變數值的設定一變就中止，避免新舊結果混在一起。
  • 開跑時自動掃描資料集（樣本數、維度、少數類數量、IR、類別欄位）→ dataset_profile.csv
"""
from __future__ import annotations

import os

# ── 必須在 import numpy / torch 之前：每個 worker 只用 1 條執行緒，平行才有效率 ──
for _v in ("OMP_NUM_THREADS", "MKL_NUM_THREADS", "OPENBLAS_NUM_THREADS",
           "NUMEXPR_NUM_THREADS", "VECLIB_MAXIMUM_THREADS"):
    os.environ.setdefault(_v, "1")
os.environ["CUDA_VISIBLE_DEVICES"] = ""          # 不用 GPU（見檔頭說明）

import re
import sys
import json
import time
import zlib
import socket
import hashlib
import argparse
import platform
import itertools
import traceback
import warnings
from dataclasses import dataclass, field, asdict
from pathlib import Path
from concurrent.futures import ProcessPoolExecutor, as_completed

import numpy as np
import pandas as pd

import sklearn
from sklearn.preprocessing import MinMaxScaler
from sklearn.svm import OneClassSVM
from sklearn.neighbors import LocalOutlierFactor
from sklearn.ensemble import IsolationForest
from sklearn.metrics import (roc_auc_score, f1_score, recall_score,
                             confusion_matrix, average_precision_score)

import torch
import torch.nn as nn
from torch.utils.data import DataLoader, TensorDataset

warnings.filterwarnings("ignore")

# ═════════════════════════ 固定設定（與 A / B / C 原程式一致）══════════════════
N_FOLDS       = 5
AE_BATCH_SIZE = 64
AE_LR         = 1e-3
DAE_NOISE     = 0.1
SAE_SPARSITY  = 1e-3
VAE_BETA      = 1.0
GLOBAL_SEED   = 42
THRESHOLD_PCT = 90          # 門檻 = 訓練多數類異常分數的第 90 百分位

N_LAYERS_LIST     = [1, 2, 3]
BOTTLENECK_RATIOS = {"1/4": 0.25, "1/3": 1/3, "1/2": 0.5, "1/1": 1.0,
                     "2/1": 2.0,  "3/1": 3.0, "4/1": 4.0}
ALL_CONFIGS = [f"h{nl}-{rl}" for nl in N_LAYERS_LIST for rl in BOTTLENECK_RATIOS]

OCC_TYPES   = ["OCSVM", "LOF", "iForest"]
METRIC_COLS = ["AUC", "F1", "Recall", "G-mean"]
EXTRA_COLS  = ["AP"]                        # Average Precision（average_precision_score），額外記錄

# ── 參數格（依文獻整理，見 00_說明 分頁）──
DEFAULT_GRID = {
    "OCSVM":   {"nu": [0.01, 0.05, 0.1, 0.2, 0.3, 0.4, 0.5],
                "gamma": ["scale", 0.01, 0.1, 1, 10, 100]},
    "LOF":     {"n_neighbors": [10, 20, 30, 40, 50]},
    "iForest": {"n_estimators": [50, 100, 150, 200],
                "max_samples": [64, 128, 256]},
}
# ── 舊程式的固定參數（一定要包含在參數格裡，才能算「調參前 vs 調參後」）──
LEGACY_DEFAULT = {
    "OCSVM":   {"nu": 0.1, "gamma": "scale"},
    "LOF":     {"n_neighbors": 20},
    "iForest": {"n_estimators": 100, "max_samples": 256},   # sklearn 'auto' = min(256, n)
}

EXCLUDE_DEFAULT = ["abalone19-5-fold", "abalone19-5-fold_Encoding"]


@dataclass
class RunConfig:
    data_root: str = "data"
    out_root: str = "results_tuning"
    run_name: str = "ABC_tune"
    n_jobs: int = 1                      # 平行 worker 數
    torch_threads_per_job: int = 1       # 每個 worker 的 torch 執行緒（平行時請保持 1）
    datasets: list = field(default_factory=list)          # 空 = 全部
    # 正式跑請填：必須「剛好」是這些資料集，多一個少一個都會停止
    expected_datasets: list = field(default_factory=list)
    exclude_datasets: list = field(default_factory=lambda: list(EXCLUDE_DEFAULT))
    studies: list = field(default_factory=lambda: ["A", "B", "C"])
    ae_types: list = field(default_factory=lambda: ["AE", "DAE", "SAE", "VAE"])
    configs: list = field(default_factory=lambda: list(ALL_CONFIGS))
    ae_epochs: int = 100
    grid: dict = field(default_factory=lambda: json.loads(json.dumps(DEFAULT_GRID)))
    # LOF 訓練集分數：
    #   "fit"    → 用 fit 時算好的訓練 LOF（negative_outlier_factor_），每點不把自己當鄰居 ← 正確做法
    #   "legacy" → 舊程式：novelty=True 還對訓練資料呼叫 decision_function（偏樂觀）
    #   只影響門檻（F1 / Recall / G-mean），不影響 AUC。
    lof_train_score: str = "fit"
    use_df_cache: bool = True
    exclude_categorical: bool = True     # 其他資料集若偵測到類別欄位也一併排除
    progress_every: int = 20             # 每完成幾格印一次進度


# ═════════════════════════ 小工具 ════════════════════════════════════════════
def stable_hash(*parts) -> int:
    return zlib.crc32("|".join(str(p) for p in parts).encode("utf-8")) & 0xFFFFFFFF


def cell_seed(ds, fold, ae, cfg_label) -> int:
    return (GLOBAL_SEED + stable_hash(ds, fold, ae, cfg_label)) % (2**31 - 1)


def safe_name(s: str) -> str:
    return re.sub(r"[^A-Za-z0-9_.-]", "-", s)


def sha1_file(p) -> str:
    return hashlib.sha1(Path(p).read_bytes()).hexdigest()[:16]


def dict_hash(d) -> str:
    return hashlib.sha256(json.dumps(d, sort_keys=True, default=str).encode()).hexdigest()[:12]


def atomic_write_csv(df: pd.DataFrame, path: Path):
    """先寫暫存檔，再一次改名成正式檔。改名是原子動作：不會出現寫一半的正式檔。"""
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(path.name + f".tmp{os.getpid()}")
    df.to_csv(tmp, index=False, encoding="utf-8")
    os.replace(tmp, path)


def param_combos(occ, grid):
    keys = list(grid[occ].keys())
    return [dict(zip(keys, vals)) for vals in itertools.product(*[grid[occ][k] for k in keys])]


def param_id(p: dict) -> str:
    return "|".join(f"{k}={v}" for k, v in p.items())


def is_legacy_default(occ, p) -> bool:
    return all(str(p.get(k)) == str(v) for k, v in LEGACY_DEFAULT[occ].items())


# ═════════════════════════ 讀資料（與 A~N 相同）══════════════════════════════
def parse_keel_dat(filepath, minority_label=None):
    lines = Path(filepath).read_text(encoding="utf-8", errors="replace").splitlines()
    data_start, rows = False, []
    for line in lines:
        s = line.strip()
        if not s or s.startswith("%"):
            continue
        if s.lower() == "@data":
            data_start = True
            continue
        if data_start:
            rows.append(s)
    if not rows:
        raise ValueError(f"No data found in {filepath}")
    df = pd.DataFrame([[p.strip() for p in r.split(",")] for r in rows])
    y_raw = df[df.columns[-1]].astype(str).str.strip().values
    feat = df.iloc[:, :-1].copy()
    n_cat = 0
    for c in feat.columns:
        conv = pd.to_numeric(feat[c], errors="coerce")
        if conv.isna().all():
            n_cat += 1
            feat[c] = pd.Categorical(feat[c]).codes.astype(float)
        else:
            feat[c] = conv
    X = feat.values.astype(float)
    if minority_label is None:
        u, cnt = np.unique(y_raw, return_counts=True)
        minority_label = u[np.argmin(cnt)]
    y = (y_raw == minority_label).astype(int)
    return X, y, minority_label, n_cat


def find_fold_files(ds_dir: Path, ds_name: str, fold: int):
    pre = re.sub(r"-fold.*$", "", ds_name)
    tra = [ds_dir / f"{pre}-{fold}tra.dat", ds_dir / f"{ds_name}-5-fold-tra{fold}.dat",
           ds_dir / f"{ds_name}-5-tra{fold}.dat", ds_dir / f"{ds_name}_fold{fold}_train.dat"]
    tst = [ds_dir / f"{pre}-{fold}tst.dat", ds_dir / f"{ds_name}-5-fold-tst{fold}.dat",
           ds_dir / f"{ds_name}-5-tst{fold}.dat", ds_dir / f"{ds_name}_fold{fold}_test.dat"]
    return (next((p for p in tra if p.exists()), None),
            next((p for p in tst if p.exists()), None))


def load_fold(data_root, ds, fold):
    """讀一折資料。任何不合格都丟出例外並說明原因（不回傳 None、不默默略過）。"""
    ds_dir = Path(data_root) / ds
    tra, tst = find_fold_files(ds_dir, ds, fold)
    if tra is None or tst is None:
        raise FileNotFoundError(f"{ds} 第 {fold} 折缺檔（train={tra}, test={tst}）")
    X_tra, y_tra, minlab, _ = parse_keel_dat(tra)
    X_tst, y_tst, _, _ = parse_keel_dat(tst, minority_label=minlab)
    if np.isnan(X_tra).any() or np.isnan(X_tst).any():
        raise ValueError(f"{ds} 第 {fold} 折有缺失值／無法轉成數字的欄位")
    if X_tra.shape[1] != X_tst.shape[1]:
        raise ValueError(f"{ds} 第 {fold} 折 train/test 欄位數不同")
    maj_rows, min_rows = np.where(y_tra == 0)[0], np.where(y_tra == 1)[0]
    if len(maj_rows) < 5:
        raise ValueError(f"{ds} 第 {fold} 折訓練多數類只有 {len(maj_rows)} 筆（<5）")
    if len(min_rows) < 1:
        raise ValueError(f"{ds} 第 {fold} 折訓練集沒有少數類")
    if len(np.unique(y_tst)) < 2:
        raise ValueError(f"{ds} 第 {fold} 折測試集只有一個類別，無法算 AUC")
    sc = MinMaxScaler()
    X_maj_s = sc.fit_transform(X_tra[maj_rows])    # fit 只用訓練多數類（與 A/B/C 相同）
    return {"X_maj_s": X_maj_s, "X_min_s": sc.transform(X_tra[min_rows]),
            "X_tst_s": sc.transform(X_tst), "y_tst": y_tst,
            "n_maj": len(maj_rows), "n_min": len(min_rows),
            "maj_rows": maj_rows, "min_rows": min_rows,
            "train_sha1": sha1_file(tra), "test_sha1": sha1_file(tst)}


# ═════════════════════════ AE（與 B/C 原程式相同的架構）═══════════════════════
class AEModel(nn.Module):
    def __init__(self, d, n_layers, n_units):
        super().__init__()
        dims = [d] + [n_units] * n_layers
        enc, dec = [], []
        for i in range(len(dims) - 1):
            enc += [nn.Linear(dims[i], dims[i + 1]), nn.ReLU()]
        for i in range(len(dims) - 1):
            act = nn.Sigmoid() if i == len(dims) - 2 else nn.ReLU()
            dec += [nn.Linear(dims[-(i + 1)], dims[-(i + 2)]), act]
        self.encoder, self.decoder = nn.Sequential(*enc), nn.Sequential(*dec)

    def forward(self, x):
        z = self.encoder(x)
        return self.decoder(z), z


class VAEModel(nn.Module):
    def __init__(self, d, n_layers, n_units):
        super().__init__()
        dims = [d] + [n_units] * n_layers
        base = []
        for i in range(len(dims) - 2):
            base += [nn.Linear(dims[i], dims[i + 1]), nn.ReLU()]
        self.enc_base = nn.Sequential(*base) if base else nn.Identity()
        mid = dims[-2] if len(dims) >= 2 else d
        self.fc_mu, self.fc_logvar = nn.Linear(mid, dims[-1]), nn.Linear(mid, dims[-1])
        dd, dec = dims[::-1], []
        for i in range(len(dd) - 1):
            act = nn.Sigmoid() if i == len(dd) - 2 else nn.ReLU()
            dec += [nn.Linear(dd[i], dd[i + 1]), act]
        self.decoder = nn.Sequential(*dec)

    def forward(self, x):
        h = self.enc_base(x)
        mu, lv = self.fc_mu(h), self.fc_logvar(h)
        z = mu + torch.exp(0.5 * lv) * torch.randn_like(lv)
        return self.decoder(z), z, mu, lv


def train_ae_extract(ae_type, X_maj_s, X_tst_s, X_min_s, n_layers, n_units, epochs, seed):
    """AE 只用訓練多數類訓練；少數類與測試集只做轉換（給之後 J/K/N 的 ENN/CNN/TL 當參考點）。"""
    torch.manual_seed(seed)
    np.random.seed(seed % (2**32 - 1))
    d = X_maj_s.shape[1]
    model = VAEModel(d, n_layers, n_units) if ae_type == "VAE" else AEModel(d, n_layers, n_units)
    opt, mse = torch.optim.Adam(model.parameters(), lr=AE_LR), nn.MSELoss()
    g = torch.Generator().manual_seed(seed)
    loader = DataLoader(TensorDataset(torch.tensor(X_maj_s, dtype=torch.float32)),
                        batch_size=min(AE_BATCH_SIZE, len(X_maj_s)), shuffle=True, generator=g)
    for _ in range(epochs):
        model.train()
        for (xb,) in loader:
            opt.zero_grad()
            if ae_type == "DAE":
                xr, _ = model(torch.clamp(xb + DAE_NOISE * torch.randn_like(xb), 0, 1))
                loss = mse(xr, xb)
            elif ae_type == "SAE":
                xr, z = model(xb)
                loss = mse(xr, xb) + SAE_SPARSITY * z.abs().mean()
            elif ae_type == "VAE":
                xr, _, mu, lv = model(xb)
                loss = mse(xr, xb) + VAE_BETA * (-0.5 * (1 + lv - mu.pow(2) - lv.exp()).mean())
            else:
                xr, _ = model(xb)
                loss = mse(xr, xb)
            loss.backward()
            opt.step()
    model.eval()

    def ext(X):
        with torch.no_grad():
            xt = torch.tensor(X, dtype=torch.float32)
            return (model(xt)[2] if ae_type == "VAE" else model(xt)[1]).numpy()
    return ext(X_maj_s), ext(X_tst_s), ext(X_min_s)


def _code_hash() -> str:
    """讀資料、前處理、AE 的程式碼指紋：改了這些程式碼，舊的 DF 快取自動失效。"""
    import inspect
    src = "".join(inspect.getsource(f) for f in
                  (parse_keel_dat, find_fold_files, load_fold, AEModel, VAEModel, train_ae_extract))
    return hashlib.sha256(src.encode()).hexdigest()[:12]


def ae_signature(cfg: RunConfig) -> dict:
    """會改變 DF 數值的設定（改 OCC 參數格不會讓 DF 快取失效）。"""
    return {"epochs": cfg.ae_epochs, "batch": AE_BATCH_SIZE, "lr": AE_LR,
            "dae_noise": DAE_NOISE, "sae_sparsity": SAE_SPARSITY, "vae_beta": VAE_BETA,
            "seed_mode": "stable_per_cell_crc32", "global_seed": GLOBAL_SEED,
            "scaler": "minmax_fit_train_majority", "code": _code_hash(),
            "torch_major_minor": ".".join(torch.__version__.split(".")[:2])}


def get_df(cfg, ds, fold, ae, cfg_label, fd):
    """先找快取；沒有或不相符才訓練 AE。回傳 (DF_maj, DF_tst, 來源)。
    快取另外存 DF_min 與原始列索引，之後 J/K/N 的欠抽樣可以直接用，不必重訓 AE。"""
    cache = (Path(cfg.out_root) / "_df_cache" / dict_hash(ae_signature(cfg)) / safe_name(ds)
             / f"f{fold}__{ae}__{safe_name(cfg_label)}.npz")
    if cfg.use_df_cache and cache.exists():
        try:
            z = np.load(cache, allow_pickle=False)
            meta = json.loads(str(z["meta"]))
            if (meta.get("train_sha1") == fd["train_sha1"] and meta.get("test_sha1") == fd["test_sha1"]
                    and "DF_min" in z.files and len(z["DF_maj"]) == fd["n_maj"]):
                return z["DF_maj"].astype(np.float64), z["DF_tst"].astype(np.float64), "cache"
        except Exception:
            pass                                          # 壞檔或舊格式就重訓
    n_layers = int(re.match(r"^h(\d+)-", cfg_label).group(1))
    ratio = BOTTLENECK_RATIOS[cfg_label.split("-", 1)[1]]
    n_units = max(2, round(fd["X_maj_s"].shape[1] * ratio))
    seed = cell_seed(ds, fold, ae, cfg_label)
    DF_maj, DF_tst, DF_min = train_ae_extract(ae, fd["X_maj_s"], fd["X_tst_s"], fd["X_min_s"],
                                              n_layers, n_units, cfg.ae_epochs, seed)
    # 一律轉成 float32 精度：確保「這次現場訓練」和「下次讀快取」算出一模一樣的結果
    DF_maj, DF_tst, DF_min = (a.astype(np.float32).astype(np.float64) for a in (DF_maj, DF_tst, DF_min))
    if cfg.use_df_cache:
        cache.parent.mkdir(parents=True, exist_ok=True)
        tmp = cache.with_name(cache.stem + f".tmp{os.getpid()}.npz")
        meta = {"dataset": ds, "fold": fold, "ae": ae, "config": cfg_label, "seed": seed,
                "n_units": n_units, "train_sha1": fd["train_sha1"], "test_sha1": fd["test_sha1"],
                "label_def": "y=1 minority(anomaly), minority label decided on the train fold",
                "ae_signature": ae_signature(cfg)}
        np.savez_compressed(tmp, DF_maj=DF_maj.astype(np.float32), DF_tst=DF_tst.astype(np.float32),
                            DF_min=DF_min.astype(np.float32), y_tst=fd["y_tst"].astype(np.int8),
                            maj_rows=fd["maj_rows"], min_rows=fd["min_rows"],
                            meta=np.array(json.dumps(meta)))
        os.replace(tmp, cache)
    return DF_maj, DF_tst, "trained"


# ═════════════════════════ OCC 評估（參數迴圈在最內層）════════════════════════
def gmean_score(y_true, y_pred):
    cm = confusion_matrix(y_true, y_pred, labels=[1, 0])
    tp, fn, fp, tn = cm[0, 0], cm[0, 1], cm[1, 0], cm[1, 1]
    sens = tp / (tp + fn) if (tp + fn) else 0.0
    spec = tn / (tn + fp) if (tn + fp) else 0.0
    return float(np.sqrt(sens * spec))


def _fit_score(occ, p, Xtr, Xte, lof_mode):
    n = len(Xtr)
    used = {}
    if occ == "OCSVM":
        clf = OneClassSVM(kernel="rbf", nu=p["nu"], gamma=p["gamma"]).fit(Xtr)
        used["gamma_used"] = float(clf._gamma)          # "scale" 實際換算出的數值（每折不同）
        s_tr, s_te = -clf.score_samples(Xtr), -clf.score_samples(Xte)
    elif occ == "LOF":
        k = int(min(p["n_neighbors"], n - 1))
        used["k_used"] = k
        clf = LocalOutlierFactor(n_neighbors=k, novelty=True).fit(Xtr)
        s_te = -clf.score_samples(Xte)
        s_tr = -clf.negative_outlier_factor_ if lof_mode == "fit" else -clf.score_samples(Xtr)
    else:
        ms = int(min(p["max_samples"], n))
        used["max_samples_used"] = ms
        clf = IsolationForest(n_estimators=p["n_estimators"], max_samples=ms,
                              random_state=GLOBAL_SEED).fit(Xtr)
        s_tr, s_te = -clf.score_samples(Xtr), -clf.score_samples(Xte)
    # 分數越高越異常。contamination 只會讓分數整體平移一個常數 → 不影響排序與門檻，所以不調。
    return s_tr, s_te, used


def eval_grid(Xtr, Xte, y_te, cfg: RunConfig, base: dict):
    """同一份特徵，跑完三個 OCC 的全部參數組合。回傳 list[dict]。"""
    rows = []
    for occ in OCC_TYPES:
        memo = {}                                   # LOF k 或 iForest max_samples 被截斷時重複利用
        for p in param_combos(occ, cfg.grid):
            row = {**base, "OCC": occ, "ParamID": param_id(p),
                   "IsDefault": is_legacy_default(occ, p), "FeatDim": int(Xtr.shape[1])}
            for k, v in p.items():
                row[f"p_{k}"] = v
            t0 = time.perf_counter()
            try:
                key = None
                if occ == "LOF":
                    key = int(min(p["n_neighbors"], len(Xtr) - 1))
                elif occ == "iForest":
                    key = (p["n_estimators"], int(min(p["max_samples"], len(Xtr))))
                if key is not None and key in memo:
                    mets, used = memo[key]
                else:
                    s_tr, s_te, used = _fit_score(occ, p, Xtr, Xte, cfg.lof_train_score)
                    thr = np.percentile(s_tr, THRESHOLD_PCT)
                    y_pred = (s_te >= thr).astype(int)
                    mets = {"AUC": roc_auc_score(y_te, s_te),
                            "F1": f1_score(y_te, y_pred, pos_label=1, zero_division=0),
                            "Recall": recall_score(y_te, y_pred, pos_label=1, zero_division=0),
                            "G-mean": gmean_score(y_te, y_pred),
                            "AP": average_precision_score(y_te, s_te)}
                    if key is not None:
                        memo[key] = (mets, used)
                row.update(mets)
                row.update(used)
                row["Status"] = "ok"
            except Exception as e:
                row.update({m: np.nan for m in METRIC_COLS + EXTRA_COLS})
                row["Status"] = f"error:{type(e).__name__}:{str(e)[:80]}"
            row["OCCSeconds"] = round(time.perf_counter() - t0, 4)
            rows.append(row)
    return rows


def n_expected_rows(cfg: RunConfig) -> int:
    return sum(len(param_combos(o, cfg.grid)) for o in OCC_TYPES)


# ═════════════════════════ 一個「格」＝ 一個平行工作單位 ══════════════════════
def cell_path(cfg: RunConfig, task) -> Path:
    root = Path(cfg.out_root) / cfg.run_name / "cells"
    if task[0] == "A":
        _, ds, fold = task
        return root / "A" / safe_name(ds) / f"fold{fold}.csv"
    _, ds, fold, ae, c = task
    return root / "DF" / safe_name(ds) / f"fold{fold}__{ae}__{safe_name(c)}.csv"


def run_task(task, cfg: RunConfig):
    """在 worker 裡執行。成功 → 寫檔並回傳統計；失敗 → 丟例外（不寫檔，下次重跑會補）。"""
    torch.set_num_threads(max(1, cfg.torch_threads_per_job))
    t0 = time.perf_counter()
    host = socket.gethostname()
    if task[0] == "A":
        _, ds, fold = task
        fd = load_fold(cfg.data_root, ds, fold)
        X_maj_s, X_tst_s, y_tst, n_maj = fd["X_maj_s"], fd["X_tst_s"], fd["y_tst"], fd["n_maj"]
        base = {"Study": "A", "Dataset": ds, "Fold": fold, "AE": "-", "Config": "-",
                "NTrainMaj": n_maj, "DFSource": "-", "Host": host}
        rows = eval_grid(X_maj_s, X_tst_s, y_tst, cfg, base)
        expected = n_expected_rows(cfg)
    else:
        _, ds, fold, ae, c = task
        fd = load_fold(cfg.data_root, ds, fold)
        X_maj_s, X_tst_s, y_tst, n_maj = fd["X_maj_s"], fd["X_tst_s"], fd["y_tst"], fd["n_maj"]
        DF_maj, DF_tst, src = get_df(cfg, ds, fold, ae, c, fd)
        rows, expected = [], 0
        common = {"Dataset": ds, "Fold": fold, "AE": ae, "Config": c,
                  "NTrainMaj": n_maj, "DFSource": src, "Host": host}
        if "B" in cfg.studies:
            sc = MinMaxScaler().fit(DF_maj)                      # 與 B 原程式相同
            rows += eval_grid(sc.transform(DF_maj), sc.transform(DF_tst), y_tst, cfg,
                              {"Study": "B", **common})
            expected += n_expected_rows(cfg)
        if "C" in cfg.studies:
            comb_maj, comb_tst = np.hstack([X_maj_s, DF_maj]), np.hstack([X_tst_s, DF_tst])
            sc = MinMaxScaler().fit(comb_maj)                    # 與 C 原程式相同
            rows += eval_grid(sc.transform(comb_maj), sc.transform(comb_tst), y_tst, cfg,
                              {"Study": "C", **common})
            expected += n_expected_rows(cfg)
    # ── 完成判定：列數正確、而且每一列都成功，才寫正式檔 ──
    if len(rows) != expected:
        raise RuntimeError(f"列數不符：{len(rows)} ≠ {expected}")
    bad = [r for r in rows if r["Status"] != "ok"]
    if bad:
        raise RuntimeError(f"{len(bad)}/{len(rows)} 組參數失敗，整格不存檔（下次重跑會再試）。"
                           f"第一個：{bad[0]['Study']} {bad[0]['OCC']} {bad[0]['ParamID']} → {bad[0]['Status']}")
    atomic_write_csv(pd.DataFrame(rows), cell_path(cfg, task))   # 寫檔成功 ＝ 這格完成
    return {"task": "|".join(map(str, task)), "kind": task[0], "dataset": task[1],
            "n_maj": int(n_maj), "seconds": round(time.perf_counter() - t0, 3),
            "df_source": rows[0].get("DFSource", "-") if rows else "-", "host": host}


def _worker_init(threads):
    torch.set_num_threads(max(1, threads))


# ═════════════════════════ 資料集掃描（第 0 步）═══════════════════════════════
def discover_datasets(cfg: RunConfig):
    root = Path(cfg.data_root)
    if not root.exists():
        raise FileNotFoundError(f"找不到資料夾 {root.resolve()}")
    all_dirs = sorted(d.name for d in root.iterdir() if d.is_dir() and not d.name.startswith("."))
    prof, keep, dropped = [], [], []
    for ds in all_dirs:
        tra, _ = find_fold_files(root / ds, ds, 1)
        if tra is None:                       # 缺檔也要留紀錄，不能默默略過
            prof.append({"Dataset": ds, "Included": False, "ExcludeReason": "找不到第 1 折 train 檔",
                         "n_categorical_cols": np.nan, "n_majority_fold1": 0, "fingerprint": "-"})
            dropped.append(ds)
            continue
        X, y, minlab, n_cat = parse_keel_dat(tra)
        n_min, n_maj = int(y.sum()), int((y == 0).sum())
        fps = []
        for f in range(1, N_FOLDS + 1):
            a, b = find_fold_files(root / ds, ds, f)
            fps += [sha1_file(a) if a else "-", sha1_file(b) if b else "-"]
        info = {"Dataset": ds, "n_train_fold1": len(y), "dim": X.shape[1],
                "n_majority_fold1": n_maj, "n_minority_fold1": n_min,
                "IR": round(n_maj / max(n_min, 1), 2), "n_categorical_cols": n_cat,
                "fingerprint": hashlib.sha1("".join(fps).encode()).hexdigest()[:16]}
        reason = ""
        if ds in cfg.exclude_datasets:
            reason = "exclude 清單"
        elif cfg.exclude_categorical and n_cat > 0:
            reason = f"含 {n_cat} 個類別欄位"
        elif cfg.datasets and ds not in cfg.datasets:
            reason = "不在本次 datasets 清單"
        info["Included"] = reason == ""
        info["ExcludeReason"] = reason
        prof.append(info)
        (keep if reason == "" else dropped).append(ds)
    prof = pd.DataFrame(prof)
    problems = []
    for d in cfg.datasets:
        if d not in all_dirs:
            problems.append(f"datasets 清單裡的「{d}」在 {root} 找不到（請檢查拼字）")
        elif d not in keep:
            r = prof[prof["Dataset"] == d].iloc[0]
            problems.append(f"datasets 清單裡的「{d}」沒有被納入：{r['ExcludeReason']}")
    if cfg.expected_datasets:
        miss = sorted(set(cfg.expected_datasets) - set(keep))
        extra = sorted(set(keep) - set(cfg.expected_datasets))
        if miss:
            problems.append(f"應納入但沒有納入：{miss}")
        if extra:
            problems.append(f"不在預定清單卻被納入：{extra}")
    if problems:
        raise SystemExit("\n✋ 資料集檢查未通過，程式停止（不會自行縮小實驗範圍）：\n   - "
                         + "\n   - ".join(problems))
    return keep, prof


def validate_folds(cfg: RunConfig, datasets):
    """開跑前把每個資料集的 5 折 train/test 全部讀一遍；任何一折有問題就停止。"""
    problems, n_files = [], 0
    for ds in datasets:
        for f in range(1, N_FOLDS + 1):
            try:
                load_fold(cfg.data_root, ds, f)
                n_files += 2
            except Exception as e:
                problems.append(str(e))
    if problems:
        raise SystemExit("\n✋ 資料檢查未通過，程式停止：\n   - " + "\n   - ".join(problems))
    print(f"資料檢查通過：{len(datasets)} 個資料集 × {N_FOLDS} 折，共 {n_files} 個檔案都可正常讀取")


# ═════════════════════════ 設定簽章（防止新舊結果混用）═══════════════════════
def hard_signature(cfg: RunConfig) -> dict:
    return {"studies": sorted(cfg.studies), "ae_types": cfg.ae_types, "configs": cfg.configs,
            "grid": cfg.grid, "legacy_default": LEGACY_DEFAULT, "ae": ae_signature(cfg),
            "lof_train_score": cfg.lof_train_score, "threshold_pct": THRESHOLD_PCT,
            "metrics": METRIC_COLS + EXTRA_COLS, "n_folds": N_FOLDS}


def check_signature(cfg: RunConfig, profile: pd.DataFrame):
    run_dir = Path(cfg.out_root) / cfg.run_name
    run_dir.mkdir(parents=True, exist_ok=True)
    sig_file = run_dir / "signature.json"
    sig = hard_signature(cfg)
    fps = dict(zip(profile["Dataset"], profile["fingerprint"]))
    if sig_file.exists():
        old = json.loads(sig_file.read_text(encoding="utf-8"))
        if old["hard"] != json.loads(json.dumps(sig, default=str)):
            diff = [k for k in sig if json.dumps(old["hard"].get(k), default=str)
                    != json.dumps(sig[k], default=str)]
            raise SystemExit(
                f"\n✋ 設定與這個 run（{cfg.run_name}）之前的設定不同：{diff}\n"
                f"   為了不讓新舊結果混在一起，程式停止。\n"
                f"   → 想保留舊結果：把 run_name 改成新名字。\n"
                f"   → 確定不要舊結果：刪掉 {run_dir} 再跑。\n"
                f"   （AE 特徵快取放在 {Path(cfg.out_root) / '_df_cache'}，只要 AE 設定沒變，新 run 仍會沿用）")
        bad = [d for d, fp in fps.items() if d in old["fingerprints"] and old["fingerprints"][d] != fp]
        if bad:
            raise SystemExit(f"\n✋ 這些資料集的檔案內容跟上次不同：{bad}\n   請改 run_name 或刪除 {run_dir}。")
        old["fingerprints"].update(fps)
        old.setdefault("sessions", []).append(soft_info(cfg))
        sig_file.write_text(json.dumps(old, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    else:
        sig_file.write_text(json.dumps({"hard": sig, "fingerprints": fps,
                                        "sessions": [soft_info(cfg)]},
                                       ensure_ascii=False, indent=2, default=str), encoding="utf-8")


def soft_info(cfg):
    return {"time": time.strftime("%Y-%m-%d %H:%M:%S"), "host": socket.gethostname(),
            "n_jobs": cfg.n_jobs, "python": platform.python_version(), "numpy": np.__version__,
            "pandas": pd.__version__, "sklearn": sklearn.__version__, "torch": torch.__version__,
            "platform": platform.platform()}


# ═════════════════════════ 主流程 ════════════════════════════════════════════
def build_tasks(cfg, datasets, n_maj_map):
    tasks = []
    if "A" in cfg.studies:
        tasks += [("A", ds, f) for ds in datasets for f in range(1, N_FOLDS + 1)]
    if {"B", "C"} & set(cfg.studies):
        tasks += [("DF", ds, f, ae, c) for ds in datasets for f in range(1, N_FOLDS + 1)
                  for ae in cfg.ae_types for c in cfg.configs]
    # 大的資料集先跑 → 最後不會剩一個大格子讓其他 worker 空等
    tasks.sort(key=lambda t: (t[0] != "A", -n_maj_map.get(t[1], 0)))
    return tasks


def cleanup_tmp(cfg):
    """清掉上次中斷留下的暫存檔：本 run 的全部清；共用的 DF 快取只清 30 分鐘以上沒動的（避免誤刪別人正在寫的）。"""
    run_dir = Path(cfg.out_root) / cfg.run_name
    cache_dir = Path(cfg.out_root) / "_df_cache"
    for base, min_age in ((run_dir, 0), (cache_dir, 1800)):
        if not base.exists():
            continue
        for p in base.rglob("*.tmp*"):
            try:
                if time.time() - p.stat().st_mtime >= min_age:
                    p.unlink()
            except Exception:
                pass


def acquire_lock(run_dir: Path):
    """同一個 run 不能同時開兩個程式。用作業系統的檔案鎖：
    程式活著就一直鎖住；程式結束或當掉，作業系統自動釋放（不會有「過期鎖」的問題）。"""
    lock_path = run_dir / "RUNNING.lock"
    fh = open(lock_path, "a+", encoding="utf-8")
    try:
        if os.name == "nt":
            import msvcrt
            fh.seek(0)
            msvcrt.locking(fh.fileno(), msvcrt.LK_NBLCK, 1)
        else:
            import fcntl
            fcntl.flock(fh.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
    except OSError:
        fh.close()
        raise SystemExit(f"\n✋ 這個 run（{run_dir.name}）已經有另一個程式在跑，不能同時開兩個。\n"
                         f"   先確認：Linux 用 `ps aux | grep run_`，Windows 看工作管理員。")
    fh.seek(0)
    fh.truncate()
    fh.write(f"host={socket.gethostname()} pid={os.getpid()} start={time.strftime('%Y-%m-%d %H:%M:%S')}")
    fh.flush()
    return fh


def release_lock(fh):
    try:
        if os.name == "nt":
            import msvcrt
            fh.seek(0)
            msvcrt.locking(fh.fileno(), msvcrt.LK_UNLCK, 1)
        fh.close()
    except Exception:
        pass


def run(cfg: RunConfig):
    run_dir = Path(cfg.out_root) / cfg.run_name
    print("=" * 76)
    print(f"OCC 調參：studies={cfg.studies}  run={cfg.run_name}  worker={cfg.n_jobs}  host={socket.gethostname()}")
    datasets, profile = discover_datasets(cfg)
    run_dir.mkdir(parents=True, exist_ok=True)
    profile.to_csv(run_dir / "dataset_profile.csv", index=False, encoding="utf-8-sig")
    print(f"資料集：納入 {len(datasets)} 個 → {datasets}")
    dropped = profile[~profile["Included"]]
    if len(dropped):
        print("未納入：" + "；".join(f"{r.Dataset}（{r.ExcludeReason}）" for r in dropped.itertuples()))
    validate_folds(cfg, datasets)
    check_signature(cfg, profile)
    lock = acquire_lock(run_dir)
    try:
        return _run_locked(cfg, run_dir, datasets, profile, lock)
    finally:
        release_lock(lock)


def _run_locked(cfg, run_dir, datasets, profile, lock):
    cleanup_tmp(cfg)
    n_maj_map = dict(zip(profile["Dataset"], profile["n_majority_fold1"]))
    tasks = build_tasks(cfg, datasets, n_maj_map)
    todo = [t for t in tasks if not cell_path(cfg, t).exists()]
    per_cell = n_expected_rows(cfg)
    print(f"參數組合：每份特徵 {per_cell} 組（OCSVM {len(param_combos('OCSVM', cfg.grid))}"
          f" / LOF {len(param_combos('LOF', cfg.grid))} / iForest {len(param_combos('iForest', cfg.grid))}）")
    print(f"工作格：共 {len(tasks)}，已完成 {len(tasks) - len(todo)}，本次要跑 {len(todo)}")
    print("=" * 76)
    if not todo:
        print("全部已完成。")
        return "done"

    timing_file = run_dir / "timing.csv"
    err_file = run_dir / "errors.log"
    done, failed, t_start = 0, 0, time.perf_counter()

    def on_result(res):
        pd.DataFrame([res]).to_csv(timing_file, mode="a", header=not timing_file.exists(),
                                   index=False, encoding="utf-8")

    def on_error(task, e):
        with open(err_file, "a", encoding="utf-8") as f:
            f.write(f"[{time.strftime('%Y-%m-%d %H:%M:%S')}] {task}: {e}\n")

    def progress():
        el = time.perf_counter() - t_start
        eta = el / max(done, 1) * (len(todo) - done - failed)
        print(f"  進度 {done + failed}/{len(todo)}（失敗 {failed}） 已用 {el/60:.1f} 分，"
              f"預估剩 {eta/60:.1f} 分", flush=True)

    try:
        if cfg.n_jobs <= 1:
            torch.set_num_threads(max(1, cfg.torch_threads_per_job))
            for t in todo:
                try:
                    on_result(run_task(t, cfg))
                    done += 1
                except Exception as e:
                    failed += 1
                    on_error(t, f"{type(e).__name__}: {e}")
                if (done + failed) % cfg.progress_every == 0:
                    progress()
        else:
            with ProcessPoolExecutor(max_workers=cfg.n_jobs, initializer=_worker_init,
                                     initargs=(cfg.torch_threads_per_job,)) as ex:
                futs = {ex.submit(run_task, t, cfg): t for t in todo}
                try:
                    for fu in as_completed(futs):
                        try:
                            on_result(fu.result())
                            done += 1
                        except Exception as e:
                            failed += 1
                            on_error(futs[fu], f"{type(e).__name__}: {e}")
                        if (done + failed) % cfg.progress_every == 0:
                            progress()
                except KeyboardInterrupt:
                    for fu in futs:
                        fu.cancel()
                    raise
    except KeyboardInterrupt:
        print("\n⏸  已中斷。已完成的格子都存好了，重新執行同一個指令就會從斷掉的地方接續。")
        return "interrupted"
    if (done + failed) % cfg.progress_every:
        progress()
    if failed:
        print(f"⚠️  {failed} 格失敗，原因見 {err_file}。重跑同一指令只會補跑這些格子。")
    remaining = [t for t in tasks if not cell_path(cfg, t).exists()]
    return "done" if not remaining else "incomplete"


# ═════════════════════════ 時間估算 ══════════════════════════════════════════
def estimate(cfg: RunConfig, target_datasets=None, target_jobs=None):
    """用已完成格子的耗時，粗估「全部資料集」要跑多久。"""
    run_dir = Path(cfg.out_root) / cfg.run_name
    tf = run_dir / "timing.csv"
    if not tf.exists():
        print("還沒有耗時資料，先跑幾個資料集。")
        return
    tm = pd.read_csv(tf)
    tm = tm[tm["df_source"] != "cache"] if (tm["df_source"] != "cache").any() else tm
    _, prof = discover_datasets(RunConfig(**{**asdict(cfg), "datasets": [], "expected_datasets": []}))
    prof = prof[prof["Included"]]                     # 全部納入的資料集（不受本次 datasets 清單限制）
    if target_datasets:
        prof = prof[prof["Dataset"].isin(target_datasets)]
    jobs = target_jobs or cfg.n_jobs
    total = 0.0
    lines = []
    for kind, per_ds in (("A", N_FOLDS), ("DF", N_FOLDS * len(cfg.ae_types) * len(cfg.configs))):
        sub = tm[tm["kind"] == kind]
        if sub.empty:
            continue
        g = sub.groupby("n_maj")["seconds"].mean()
        if len(g) >= 2:                               # 以二次式配適（OCSVM 約與 n² 成正比）
            deg = 2 if len(g) >= 3 else 1
            coef = np.polyfit(g.index.values.astype(float), g.values, deg)
            f = lambda n: max(float(np.polyval(coef, n)), float(g.min()) * 0.5)
        else:
            f = lambda n: float(g.iloc[0]) * max(n / g.index[0], 0.5)
        sec = sum(f(n) * per_ds for n in prof["n_majority_fold1"])
        total += sec
        lines.append(f"  {kind:>2} 格：單核合計約 {_hm(sec)}")
    eff = 0.75 if jobs > 1 else 1.0                   # 平行效率打個折（超執行緒、I/O）
    print("── 全跑時間粗估（依已完成格子外推，僅供參考）──")
    print("\n".join(lines))
    print(f"  {len(prof)} 個資料集，單核合計約 {_hm(total)} → {jobs} 個 worker 約 {_hm(total/(jobs*eff))}")
    print("  （估算用的是「這台機器」測到的速度；換機器請用 --estimate-jobs 換算，或在新機器上先跑 2 個資料集）")


def _hm(sec):
    return f"{sec/3600:.1f} 小時" if sec >= 3600 else f"{sec/60:.0f} 分鐘"


# ═════════════════════════ 統整與 Excel ═════════════════════════════════════
def load_all(cfg: RunConfig) -> pd.DataFrame:
    root = Path(cfg.out_root) / cfg.run_name / "cells"
    files = sorted(root.rglob("*.csv"))
    if not files:
        return pd.DataFrame()
    return pd.concat((pd.read_csv(f) for f in files), ignore_index=True)


def _fmt(m, s):
    return "" if pd.isna(m) else f"{m:.4f} ± {s:.4f}" if not pd.isna(s) else f"{m:.4f}"


def select_rows(df, G, cand, only_default=False):
    """每組 G（例如 Study×Dataset×AE×OCC）挑 5-fold 平均 AUC 最高的候選，回傳該候選的所有 fold 列。
    只考慮「第 1~5 折全部都有、AUC 都有效、沒有重複」的候選；
    同分時優先選舊預設參數，其次依名稱排序（結果可重現）。"""
    d = df[(df["Status"] == "ok") & df["AUC"].notna()]
    if only_default:
        d = d[d["IsDefault"]]
    agg = d.groupby(G + cand).agg(mAUC=("AUC", "mean"), nF=("Fold", "nunique"), nRows=("Fold", "size"),
                                  isdef=("IsDefault", "first")).reset_index()
    agg = agg[(agg["nF"] == N_FOLDS) & (agg["nRows"] == N_FOLDS)]
    agg["同分候選數"] = agg.groupby(G)["mAUC"].transform(lambda x: int((x >= x.max() - 1e-12).sum()))
    agg = agg.sort_values(G + ["mAUC", "isdef"] + cand, ascending=[True] * len(G) + [False, False] + [True] * len(cand))
    pick = agg.drop_duplicates(G)[G + cand + ["同分候選數"]]
    return d.merge(pick, on=G + cand)


def overall_table(rows, label):
    keys = ["Study", "AE", "OCC"]
    out = []
    for k, g in rows.groupby(keys, sort=False):
        r = dict(zip(keys, k))
        r["選法"] = label
        r["資料集數"] = g["Dataset"].nunique()
        r["fold 數"] = len(g)
        for m in METRIC_COLS + EXTRA_COLS:
            r[f"{m}_mean"], r[f"{m}_std"] = g[m].mean(), g[m].std(ddof=1)
            r[m] = _fmt(r[f"{m}_mean"], r[f"{m}_std"])
        per_ds = g.drop_duplicates("Dataset")
        r["最常選中參數"] = per_ds["ParamID"].mode().iloc[0]
        r["最常選中 Config"] = per_ds["Config"].mode().iloc[0]
        out.append(r)
    return pd.DataFrame(out)


def summarize(cfg: RunConfig):
    run_dir = Path(cfg.out_root) / cfg.run_name
    planned, _ = discover_datasets(cfg)
    df = load_all(cfg)
    if df.empty:
        print("沒有結果可統整。")
        return None
    df = df[df["Dataset"].isin(planned) & df["Study"].isin(cfg.studies)].copy()
    df["IsDefault"] = df["IsDefault"].astype(str).str.lower().eq("true")
    for c in ["p_gamma"]:
        if c in df:
            df[c] = df[c].astype(str)
    df.to_csv(run_dir / "all_results.csv.gz", index=False, compression="gzip")

    G = ["Study", "Dataset", "AE", "OCC"]
    cand = ["Config", "ParamID"]
    best = select_rows(df, G, cand)
    dflt = select_rows(df, G, cand, only_default=True)
    if best.empty or dflt.empty:
        log = build_log(cfg, df, planned)
        out = run_dir / f"{cfg.run_name}_summary_PARTIAL.xlsx"
        write_excel(out, [("00_說明", readme_table(cfg, df, log, False, 0), None), ("06_執行紀錄", log, None)])
        print(f"⚠️  目前還沒有任何設定跑滿 5 折，只輸出進度：{out.resolve()}")
        return out

    ov_best = overall_table(best, "AE config 與 OCC 參數聯合搜尋的最佳組合")
    ov_dflt = overall_table(dflt, "原固定 OCC 參數下的最佳 AE config（＝舊表做法）")
    gain = ov_best[["Study", "AE", "OCC", "資料集數", "AUC", "AUC_mean", "最常選中參數", "最常選中 Config"]].merge(
        ov_dflt[["Study", "AE", "OCC", "AUC", "AUC_mean"]], on=["Study", "AE", "OCC"],
        suffixes=("_聯合搜尋", "_原固定參數"))
    gain["ΔAUC（聯合搜尋−原固定參數）"] = gain["AUC_mean_聯合搜尋"] - gain["AUC_mean_原固定參數"]
    order = {"A": 0, "B": 1, "C": 2}
    gain = gain.sort_values(["Study", "OCC", "AE"], key=lambda s: s.map(order) if s.name == "Study" else s)
    gain = gain[["Study", "AE", "OCC", "資料集數", "AUC_聯合搜尋", "AUC_原固定參數",
                 "ΔAUC（聯合搜尋−原固定參數）", "最常選中參數", "最常選中 Config"]]

    per_ds, fold_detail = best_param_tables(best, dflt)

    # 全域固定參數排名：同一組參數用在所有資料集（B/C 仍每資料集挑 config）→ 平均
    fixed = select_rows(df, G + ["ParamID"], ["Config"])
    fx = (fixed.groupby(["Study", "AE", "OCC", "ParamID"])
               .agg(AUC_mean=("AUC", "mean"), AUC_std=("AUC", "std"),
                    資料集數=("Dataset", "nunique"), 是否舊預設=("IsDefault", "first")).reset_index())
    fx["排名"] = fx.groupby(["Study", "AE", "OCC"])["AUC_mean"].rank(ascending=False, method="min").astype(int)
    fx = fx.sort_values(["Study", "AE", "OCC", "排名"])

    log = build_log(cfg, df, planned)
    complete = bool(log["完整"].all()) and bool((log["結果列數"] == log["應有結果列數"]).all())
    dup = int(df.duplicated(["Study", "Dataset", "AE", "Config", "Fold", "OCC", "ParamID"]).sum())
    if dup:
        complete = False

    readme = readme_table(cfg, df, log, complete, dup)
    out = run_dir / (f"{cfg.run_name}_summary.xlsx" if complete else f"{cfg.run_name}_summary_PARTIAL.xlsx")
    write_excel(out, [
        ("00_說明", readme, None),
        ("01_聯合搜尋vs原固定參數", gain, ["ΔAUC（聯合搜尋−原固定參數）"]),
        ("02_總表_聯合搜尋", ov_best.drop(columns=[c for c in ov_best if c.endswith(("_mean", "_std"))]), None),
        ("03_總表_原固定參數", ov_dflt.drop(columns=[c for c in ov_dflt if c.endswith(("_mean", "_std"))]), None),
        ("04_每資料集最佳參數", per_ds, ["ΔAUC（vs 原固定參數）"]),
        ("04b_最佳組合逐折明細", fold_detail, None),
        ("05_固定參數排名", fx, None),
        ("06_執行紀錄", log, None),
    ])
    print(f"✅ 統整完成：{out.resolve()}")
    print(f"   全部原始結果：{(run_dir / 'all_results.csv.gz').resolve()}")
    incomplete = log[~log["完整"]]
    if complete:
        print(f"✔ 驗收通過：{len(planned)} 個資料集 × {cfg.studies} 全部完成，"
              f"共 {int(log['結果列數'].sum()):,} 列結果，無重複。")
    else:
        print(f"⚠️  尚未完成（{len(incomplete)} 個 資料集×Study 缺格；重複列 {dup}）→ 檔名標 _PARTIAL，"
              f"只能當進度參考，不可用於論文。")
    print("\n── 調參 vs 舊預設（AUC）──")
    print(gain[["Study", "AE", "OCC", "AUC_聯合搜尋", "AUC_原固定參數", "ΔAUC（聯合搜尋−原固定參數）"]].to_string(index=False))
    return out


PARAM_COLS = {"OCSVM": ["nu", "gamma"], "LOF": ["n_neighbors"], "iForest": ["n_estimators", "max_samples"]}


def best_param_tables(best, dflt):
    """04：每個 (方法, 資料集, AE, OCC) 的最佳 (config, 參數)，參數拆成獨立欄位；
       04b：選中組合的 5 折逐折明細（含實際生效的參數值）。"""
    G = ["Study", "Dataset", "AE", "OCC"]
    rows = []
    for key, g in best.groupby(G, sort=False):
        r = dict(zip(G, key))
        r["最佳 Config"] = g["Config"].iloc[0]
        r["特徵維度"] = int(g["FeatDim"].iloc[0])
        for occ, ps in PARAM_COLS.items():                    # 指定的參數值；非此 OCC 的欄位留空
            for pn in ps:
                r[pn] = g[f"p_{pn}"].iloc[0] if (r["OCC"] == occ and f"p_{pn}" in g) else None
        if r["OCC"] == "LOF" and "k_used" in g:             # 實際生效值（被截斷或 scale 換算），5 折範圍
            r["實際生效"] = f"k_used {int(g['k_used'].min())}~{int(g['k_used'].max())}"
        elif r["OCC"] == "iForest" and "max_samples_used" in g:
            r["實際生效"] = f"max_samples_used {int(g['max_samples_used'].min())}~{int(g['max_samples_used'].max())}"
        elif r["OCC"] == "OCSVM" and "gamma_used" in g:
            r["實際生效"] = f"gamma_used {g['gamma_used'].min():.4g}~{g['gamma_used'].max():.4g}"
        else:
            r["實際生效"] = ""
        r["ParamID"] = g["ParamID"].iloc[0]
        r["是否＝原固定參數"] = bool(g["IsDefault"].iloc[0])
        r["AUC 平均"], r["AUC 標準差"] = g["AUC"].mean(), g["AUC"].std(ddof=1)
        for m in ["F1", "Recall", "G-mean"] + EXTRA_COLS:
            r[m] = g[m].mean()
        r["有效 folds"] = int(g["Fold"].nunique())
        r["同分候選數"] = int(g["同分候選數"].iloc[0])
        rows.append(r)
    per = pd.DataFrame(rows)
    d_auc = dflt.groupby(G)["AUC"].mean().rename("AUC（原固定參數）").reset_index()
    per = per.merge(d_auc, on=G, how="left")
    per["ΔAUC（vs 原固定參數）"] = per["AUC 平均"] - per["AUC（原固定參數）"]
    per["選取準則"] = "5 折平均 AUC 最高；F1/Recall/G-mean/AP 皆為同一組合的值"
    per["選取性質"] = "Exploratory／oracle（以測試折挑選）"
    per = per.sort_values(["Study", "OCC", "AE", "Dataset"]).reset_index(drop=True)
    keep = ["Study", "Dataset", "AE", "OCC", "Config", "ParamID", "Fold", "FeatDim",
            "NTrainMaj", "k_used", "max_samples_used", "gamma_used"] + METRIC_COLS + EXTRA_COLS
    detail = best[[c for c in keep if c in best.columns]].sort_values(
        ["Study", "OCC", "AE", "Dataset", "Fold"]).reset_index(drop=True)
    return per, detail


def build_log(cfg, df, planned):
    """執行紀錄：從「預定的完整任務清單」反查，沒產生任何結果的資料集也會列出來。"""
    rows = []
    for ds in planned:
        for st in cfg.studies:
            sub = df[(df["Dataset"] == ds) & (df["Study"] == st)]
            cells = sub.groupby(["Fold", "AE", "Config"])["DFSource"].first() if len(sub) else pd.Series(dtype=str)
            exp = N_FOLDS * (1 if st == "A" else len(cfg.ae_types) * len(cfg.configs))
            rows.append({"Dataset": ds, "Study": st, "完成格數": len(cells), "應有格數": exp,
                         "完整": len(cells) == exp, "結果列數": len(sub),
                         "應有結果列數": exp * n_expected_rows(cfg),
                         "DF 讀自快取的格數": int((cells == "cache").sum()),
                         "主機": ",".join(sorted(sub["Host"].astype(str).unique())) if len(sub) else "-"})
    return pd.DataFrame(rows)


def readme_table(cfg, df, log, complete=True, dup=0):
    g = cfg.grid
    items = [
        ("結果狀態", "✔ 正式：預定的資料集與 Study 全部完成、無重複"
         if complete else "⚠ 暫存（PARTIAL）：尚未全部完成，只能看進度，不可用於論文"),
        ("這份檔案是什麼", f"Baseline A/B/C 的 OCC 調參結果（run = {cfg.run_name}）"),
        ("資料集", f"{df['Dataset'].nunique()} 個：{', '.join(sorted(df['Dataset'].unique()))}"),
        ("排除", ", ".join(cfg.exclude_datasets) + "（abalone19 兩個版本：名目欄位 Sex 被編成 0/1/2，"
                  "且 train/test 各自編碼；會對距離型方法引入人為順序）"),
        ("OCSVM 參數格", f"nu ∈ {g['OCSVM']['nu']}；gamma ∈ {g['OCSVM']['gamma']}；kernel = rbf"),
        ("LOF 參數格", f"n_neighbors ∈ {g['LOF']['n_neighbors']}（超過 訓練樣本數−1 時自動截斷，見 k_used 欄）；novelty = True"),
        ("iForest 參數格", f"n_estimators ∈ {g['iForest']['n_estimators']}；max_samples ∈ {g['iForest']['max_samples']}"
                           "（超過訓練樣本數時自動截斷）；random_state = 42"),
        ("為什麼不調 contamination", "contamination 只讓異常分數整體平移一個常數：AUC 只看排序不受影響；"
                                   "門檻是訓練分數的第 90 百分位，也跟著平移同一個常數，所以 F1/Recall/G-mean 也不變。"),
        ("舊預設參數", "OCSVM nu=0.1, gamma=scale｜LOF k=20｜iForest 100 棵, max_samples=256(=sklearn auto)。"
                     "舊預設就在參數格裡，所以同一次執行就能比較「調參前後」。"),
        ("選法（01/02/04 分頁）", "每個 (Study, 資料集, AE, OCC) 挑 5-fold 平均 AUC 最高的 (config, 參數)；"
                               "只考慮 5 折全部完整的候選。總表是把挑中的 5 折 × 資料集全部 fold 平均 ± 標準差"
                               "（標準差跨全部 fold；與週報算法相同，5 折完整時平均值等同資料集等權平均）。"
                               "⚠ 用測試集挑選 → 屬探索性上界（oracle），論文需註明。"),
        ("05 固定參數排名", "同一組 OCC 參數套用到所有資料集的平均 AUC 排名，用來看「參數敏感度」與舊預設排第幾。"
                          "⚠ 這不是獨立驗證：B/C 的 AE config 仍逐資料集用測試結果挑，排名本身也來自同一批測試成績。"),
        ("參數依據：直接沿用", "葉芷妍 (2026) 碩論 表3-6（PDF 第24頁、印刷頁23）：OCSVM nu = 0.01/0.05/0.1/0.2/0.3/0.4/0.5、"
                             "gamma = scale｜LOF n_neighbors = 10/20/30/40/50｜iForest n_estimators = 50/100/150/200。"),
        ("參數依據：本研究擴充", "① OCSVM gamma 另加 0.01/0.1/1/10/100（對數格點）：參考 Wang et al. 2018 §4 以對數尺度搜尋核寬度；"
                              "注意該文搜尋的是 σ，若核函數為 exp(-‖x−y‖²/(2σ²)) 則 γ = 1/(2σ²)，數值不能直接照搬，"
                              "本研究的 gamma 值為自行設定。② iForest max_samples = 64/128/256：Liu et al. 2012 §4.1 建議 ψ=256、"
                              "§5.5 顯示小 ψ 即收斂，故往下擴充。"),
        ("參數依據：範圍參考", "Goldstein & Uchida 2016：k 取 10~50、nu 取 0.2~0.8（該文是「平均」不同參數的結果，不是挑最佳）｜"
                            "Breunig et al. 2000 §6.2：k 至少 10｜Schölkopf et al. 2001 §5 Prop.4：nu ∈ (0,1]｜"
                            "Xu et al. 2023 §6.1.3：iForest 常用 ψ=256。"),
        ("選取單位", "每個 (方法, 資料集, AE 類型, OCC) 各自挑一組；B/C 在 21 個 AE config × OCC 參數中聯合挑選"
                    "（例：B/glass0/VAE/OCSVM 從 21×42＝882 組中挑），A 只挑 OCC 參數。不會在 4 種 AE 之間再挑總冠軍。"),
        ("原固定參數 vs 聯合搜尋", "兩邊都會重新挑 AE config，所以差異同時包含「架構改選」與「OCC 調參」，不能全部歸因於 OCC 調參。"),
        ("C 的意思", "C 是 OF 與 DF「欄位串接」，不是數值相加。21 個 config 含 1/1~4/1，不全是降維，應寫「深度特徵萃取」。"),
        ("AP", "AP = average_precision_score（平均精確率），不是梯形積分的 PR 曲線面積。"),
        ("門檻", "訓練多數類異常分數的第 90 百分位。LOF 的訓練分數："
                + ("用 fit 時算好的 LOF（每點不把自己當鄰居，正確做法）" if cfg.lof_train_score == "fit"
                   else "舊做法（對訓練資料呼叫 decision_function）")
                + "。只影響 F1/Recall/G-mean，不影響 AUC。"),
        ("前處理", "A：MinMax(fit 訓練多數類)｜B：AE 特徵再 MinMax(fit DF_maj)｜C：[OF,DF] 串接後 MinMax。與原程式相同。"),
        ("AE", f"{cfg.ae_types}｜21 configs｜epochs={cfg.ae_epochs}, batch={AE_BATCH_SIZE}, lr={AE_LR}"),
        ("亂數", "每個 (資料集, fold, AE, config) 由 CRC32 產生固定種子 → 與執行順序、平行數量無關，"
                "但數值不會等於舊 B/C（舊程式只在開頭設一次種子）。"),
        ("運算", "全程 CPU，不使用 GPU。"),
        ("完整性", f"{int(log['完整'].sum())}/{len(log)} 個（資料集×Study）完整；"
                  f"結果列 {int(log['結果列數'].sum()):,} / 應有 {int(log['應有結果列數'].sum()):,}；重複列 {dup}。"
                  "失敗的格子不會存檔（見 errors.log），所以這裡的每一列都是成功結果。"),
        ("參考文獻", "Schölkopf et al. (2001) Neural Computation 13(7)｜Wang et al. (2018) Pattern Recognition 74｜"
                    "Goldstein & Uchida (2016) PLOS ONE 11(4)｜Breunig et al. (2000) SIGMOD｜"
                    "Campos et al. (2016) DMKD 30(4)｜Liu, Ting & Zhou (2012) ACM TKDD 6(1)｜"
                    "Xu et al. (2023) IEEE TKDE 35(12)｜葉芷妍 (2026) 中央大學碩論 表3-6（頁碼請以正式版再確認）"),
        ("版本", json.dumps(soft_info(cfg), ensure_ascii=False)),
    ]
    return pd.DataFrame(items, columns=["項目", "內容"])


def write_excel(path, sheets):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    hf = PatternFill("solid", fgColor="2F5597")
    hfont = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    bfont = Font(name="Arial", size=10)
    pos, neg = PatternFill("solid", fgColor="C6EFCE"), PatternFill("solid", fgColor="F8CBAD")
    thin = Border(*(Side(style="thin"),) * 4)
    wb = Workbook()
    wb.remove(wb.active)
    for name, df, delta_cols in sheets:
        ws = wb.create_sheet(name)
        cols = list(df.columns)
        for j, c in enumerate(cols, 1):
            cell = ws.cell(1, j, c)
            cell.font, cell.fill, cell.border = hfont, hf, thin
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for i, row in enumerate(df.itertuples(index=False), 2):
            for j, v in enumerate(row, 1):
                if isinstance(v, (np.floating, float)) and np.isnan(v):
                    v = None
                elif isinstance(v, np.generic):
                    v = v.item()
                cell = ws.cell(i, j, v)
                cell.font, cell.border = bfont, thin
                if isinstance(v, float):
                    cell.number_format = "0.0000"
                if delta_cols and cols[j - 1] in delta_cols and isinstance(v, (int, float)):
                    cell.fill = pos if v > 1e-9 else neg if v < -1e-9 else PatternFill()
        for j, c in enumerate(cols, 1):
            w = max([len(str(c))] + [len(str(x)) for x in df[c].head(200)]) if len(df) else len(str(c))
            ws.column_dimensions[get_column_letter(j)].width = min(max(10, w * 1.1 + 2), 110 if name == "00_說明" else 40)
        ws.freeze_panes = "A2"
    wb.save(path)


# ═════════════════════════ 入口 ══════════════════════════════════════════════
def main(cfg: RunConfig, argv=None):
    ap = argparse.ArgumentParser(description="A/B/C OCC 調參")
    ap.add_argument("--summarize-only", action="store_true", help="不跑實驗，只把已完成的結果統整成 Excel")
    ap.add_argument("--estimate", action="store_true", help="依已完成格子估算全部資料集要跑多久")
    ap.add_argument("--estimate-jobs", type=int, default=None, help="估算時假設的 worker 數（例如 Azure 的 8）")
    args = ap.parse_args(argv)
    if args.estimate:
        estimate(cfg, target_jobs=args.estimate_jobs)
        return
    if not args.summarize_only:
        status = run(cfg)
        if status == "interrupted":
            return
    summarize(cfg)
    if not args.summarize_only:
        estimate(cfg, target_jobs=args.estimate_jobs)
